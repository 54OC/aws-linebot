from flask import Flask, request
import requests
import os
import random
import re
from urllib.parse import parse_qs

app = Flask(__name__)

# LINE 設定
CHANNEL_ACCESS_TOKEN = os.environ.get('CHANNEL_ACCESS_TOKEN')
CHANNEL_SECRET = os.environ.get('CHANNEL_SECRET')
LINE_REPLY_URL = 'https://api.line.me/v2/bot/message/reply'

# 莫蘭迪配色方案
COLORS = {
    'primary': '#8B9D83',      # 莫蘭迪綠
    'secondary': '#B4A7A3',    # 莫蘭迪米
    'accent': '#9E9FA5',       # 莫蘭迪灰藍
    'success': '#8FA87F',      # 莫蘭迪淺綠
    'error': '#C8A8A0',        # 莫蘭迪粉
    'saa': '#7C9885',          # SAA 專用色
    'sap': '#A98E8F',          # SAP 專用色
}

# 全域題庫
QUESTIONS = {'SAA': [], 'SAP': []}

def load_questions_from_excel():
    """從 Excel 載入題庫（強化版）"""
    filename = 'AWS_SAA_SAP_繁體中文版.xlsx'
    
    print(f"🔍 開始載入 Excel：{filename}")
    
    # 檢查檔案是否存在
    if not os.path.exists(filename):
        print(f"❌ 檔案不存在：{filename}")
        print(f"📂 當前目錄檔案：{os.listdir('.')}")
        return load_default_questions()
    
    try:
        import openpyxl
        print("✅ openpyxl 套件載入成功")
        
        # 載入 Excel
        wb = openpyxl.load_workbook(filename, data_only=True)
        print(f"✅ Excel 檔案開啟成功")
        print(f"📋 工作表列表：{wb.sheetnames}")
        
        # 處理每個工作表
        for sheet_name in wb.sheetnames:
            # 只處理 SAA 和 SAP（忽略大小寫）
            exam_type = sheet_name.upper()
            if exam_type not in ['SAA', 'SAP']:
                print(f"⚠️ 跳過工作表：{sheet_name}")
                continue
            
            print(f"\n{'='*50}")
            print(f"📖 處理工作表：{sheet_name}")
            print(f"{'='*50}")
            
            sheet = wb[sheet_name]
            questions = []
            
            # 讀取標題列
            headers = []
            for cell in sheet[1]:
                if cell.value:
                    headers.append(str(cell.value).strip())
                else:
                    headers.append('')
            
            print(f"欄位列表：{headers}")
            
            # 建立欄位索引映射（支援多種欄位名稱）
            col_map = {}
            for idx, header in enumerate(headers):
                h = header.lower()
                # 題目欄位
                if '題目' in h and '中' in h:
                    col_map['q_cn'] = idx
                elif '題目' in h and '英' in h:
                    col_map['q_en'] = idx
                # 選項欄位
                elif '選項' in h and '中' in h:
                    col_map['opt_cn'] = idx
                elif '選項' in h and '英' in h:
                    col_map['opt_en'] = idx
                # 答案欄位
                elif '答案' in h or 'answer' in h:
                    col_map['ans'] = idx
                # 解析欄位
                elif ('解析' in h or 'explanation' in h) and '中' in h:
                    col_map['exp_cn'] = idx
                elif ('解析' in h or 'explanation' in h) and '英' in h:
                    col_map['exp_en'] = idx
                elif '專業解析' in h:
                    col_map['exp_cn'] = idx
            
            print(f"欄位映射：{col_map}")
            
            # 如果沒找到必要欄位，跳過這個工作表
            if 'ans' not in col_map:
                print(f"❌ 找不到「答案」欄位，跳過此工作表")
                continue
            
            # 讀取資料（從第2列開始）
            success_count = 0
            for row_idx in range(2, sheet.max_row + 1):
                try:
                    row = list(sheet[row_idx])
                    
                    # 提取資料
                    q_data = {
                        'q_cn': str(row[col_map.get('q_cn', 0)].value or '').strip() if 'q_cn' in col_map else '',
                        'q_en': str(row[col_map.get('q_en', 1)].value or '').strip() if 'q_en' in col_map else '',
                        'opt_cn_raw': str(row[col_map.get('opt_cn', 2)].value or '').strip() if 'opt_cn' in col_map else '',
                        'opt_en_raw': str(row[col_map.get('opt_en', 3)].value or '').strip() if 'opt_en' in col_map else '',
                        'ans': str(row[col_map['ans']].value or 'A').strip().upper()[0],  # 只取第一個字元
                        'exp_cn': str(row[col_map.get('exp_cn', 5)].value or '暫無中文解析').strip() if 'exp_cn' in col_map else '暫無中文解析',
                        'exp_en': str(row[col_map.get('exp_en', 6)].value or 'No explanation').strip() if 'exp_en' in col_map else 'No explanation'
                    }
                    
                    # 解析選項
                    q_data['opts_cn'] = parse_options(q_data['opt_cn_raw'])
                    q_data['opts_en'] = parse_options(q_data['opt_en_raw'])
                    
                    # 驗證：至少要有題目和選項
                    has_valid_cn = q_data['q_cn'] and len(q_data['opts_cn']) >= 4
                    has_valid_en = q_data['q_en'] and len(q_data['opts_en']) >= 4
                    
                    if has_valid_cn or has_valid_en:
                        questions.append(q_data)
                        success_count += 1
                    
                except Exception as e:
                    print(f"⚠️ 第 {row_idx} 列解析失敗：{e}")
                    continue
            
            QUESTIONS[exam_type] = questions
            print(f"✅ {exam_type}：成功載入 {success_count} / {sheet.max_row - 1} 題")
        
        wb.close()
        
        # 如果兩個都是空的
        total_questions = len(QUESTIONS['SAA']) + len(QUESTIONS['SAP'])
        if total_questions == 0:
            print("❌ Excel 解析後沒有任何題目，使用預設題庫")
            return load_default_questions()
        
        print(f"\n🎉 題庫載入完成！")
        print(f"📊 SAA: {len(QUESTIONS['SAA'])} 題")
        print(f"📊 SAP: {len(QUESTIONS['SAP'])} 題")
        
        return QUESTIONS
        
    except ImportError:
        print("❌ 缺少 openpyxl 套件")
        return load_default_questions()
    except Exception as e:
        print(f"❌ Excel 載入失敗：{type(e).__name__}: {e}")
        import traceback
        traceback.print_exc()
        return load_default_questions()

def parse_options(opt_raw):
    """解析選項（強化版）"""
    opts = {}
    
    if not opt_raw or opt_raw == 'None':
        return opts
    
    # 多種格式支援
    patterns = [
        r'A[.、:：)\s]+([^\n]+?)(?=\s*B[.、:：)\s]+|$)',
        r'B[.、:：)\s]+([^\n]+?)(?=\s*C[.、:：)\s]+|$)',
        r'C[.、:：)\s]+([^\n]+?)(?=\s*D[.、:：)\s]+|$)',
        r'D[.、:：)\s]+([^\n]+?)(?=\s*$)',
    ]
    
    for i, pattern in enumerate(patterns):
        char = ['A', 'B', 'C', 'D'][i]
        match = re.search(pattern, opt_raw, re.DOTALL | re.IGNORECASE)
        if match:
            opts[char] = match.group(1).strip()
        else:
            opts[char] = ''
    
    return opts

def load_default_questions():
    """預設題庫"""
    print("📚 載入預設題庫")
    return {
        'SAA': [
            {
                'q_cn': '某公司需要在AWS上部署一個需要處理突發流量的網站，但預算有限。以下哪種EC2購買選項最適合？',
                'q_en': 'A company needs to deploy a website on AWS that can handle traffic spikes with a limited budget. Which EC2 purchasing option is most suitable?',
                'opts_cn': {
                    'A': 'On-Demand Instances',
                    'B': 'Reserved Instances',
                    'C': 'Spot Instances',
                    'D': 'Savings Plans'
                },
                'opts_en': {
                    'A': 'On-Demand Instances',
                    'B': 'Reserved Instances',
                    'C': 'Spot Instances',
                    'D': 'Savings Plans'
                },
                'ans': 'C',
                'exp_cn': 'Spot Instances 可以以更低的價格使用未使用的 EC2 容量，最適合突發流量且預算有限的場景。',
                'exp_en': 'Spot Instances allow you to use spare EC2 capacity at a lower price.'
            },
            {
                'q_cn': '企業應用程序需要共享文件系統，可以同時被多個EC2實例訪問和修改。以下哪個AWS服務最適合？',
                'q_en': 'An enterprise application needs a shared file system that can be accessed by multiple EC2 instances. Which AWS service is most suitable?',
                'opts_cn': {
                    'A': 'Amazon S3',
                    'B': 'Amazon EBS',
                    'C': 'Amazon EFS',
                    'D': 'AWS Storage Gateway'
                },
                'opts_en': {
                    'A': 'Amazon S3',
                    'B': 'Amazon EBS',
                    'C': 'Amazon EFS',
                    'D': 'AWS Storage Gateway'
                },
                'ans': 'C',
                'exp_cn': 'Amazon EFS 提供完全託管的共享文件系統，支援多個 EC2 實例同時讀寫。',
                'exp_en': 'Amazon EFS provides a fully managed shared file system.'
            }
        ],
        'SAP': [
            {
                'q_cn': '企業需要設計一個跨多個AWS區域的災難恢復方案，RTO要求小於1小時。最適合的策略是？',
                'q_en': 'An enterprise needs a DR solution across multiple AWS regions with RTO < 1 hour. What is the most suitable strategy?',
                'opts_cn': {
                    'A': 'Backup and Restore',
                    'B': 'Pilot Light',
                    'C': 'Warm Standby',
                    'D': 'Multi-Site Active-Active'
                },
                'opts_en': {
                    'A': 'Backup and Restore',
                    'B': 'Pilot Light',
                    'C': 'Warm Standby',
                    'D': 'Multi-Site Active-Active'
                },
                'ans': 'C',
                'exp_cn': 'Warm Standby 策略在備援區域維持縮小版本的完整環境，可以快速擴展以滿足 RTO < 1小時的要求。',
                'exp_en': 'Warm Standby maintains a scaled-down environment that can be quickly scaled up.'
            }
        ]
    }

# 初始化
print("\n" + "="*60)
print("🚀 AWS LINE Bot 啟動中...")
print("="*60)
load_questions_from_excel()
print("="*60 + "\n")

def get_random_q(exam_type, lang):
    """隨機抽題"""
    pool = QUESTIONS.get(exam_type, [])
    if not pool:
        return None
    
    if lang == 'cn':
        valid = [i for i, q in enumerate(pool) if q['q_cn'] and len(q.get('opts_cn', {})) >= 4]
    else:
        valid = [i for i, q in enumerate(pool) if q['q_en'] and len(q.get('opts_en', {})) >= 4]
    
    return random.choice(valid) if valid else None

def send_line(reply_token, messages):
    """發送訊息"""
    headers = {
        'Content-Type': 'application/json',
        'Authorization': f'Bearer {CHANNEL_ACCESS_TOKEN}'
    }
    payload = {'replyToken': reply_token, 'messages': messages}
    response = requests.post(LINE_REPLY_URL, headers=headers, json=payload)
    if response.status_code != 200:
        print(f"❌ LINE API Error: {response.status_code}")

# Flex Messages（莫蘭迪配色）
def welcome_flex():
    """歡迎畫面"""
    return {
        "type": "flex",
        "altText": "AWS 證照助手",
        "contents": {
            "type": "bubble",
            "header": {
                "type": "box",
                "layout": "vertical",
                "contents": [{
                    "type": "text",
                    "text": "AWS 證照練習助手",
                    "weight": "bold",
                    "color": COLORS['primary'],
                    "size": "xl",
                    "align": "center"
                }],
                "backgroundColor": "#F5F5F0"
            },
            "body": {
                "type": "box",
                "layout": "vertical",
                "contents": [
                    {
                        "type": "text",
                        "text": "準備好挑戰了嗎？",
                        "weight": "bold",
                        "size": "md",
                        "color": "#5A5A5A"
                    },
                    {
                        "type": "box",
                        "layout": "vertical",
                        "margin": "lg",
                        "spacing": "sm",
                        "contents": [
                            {
                                "type": "text",
                                "text": f"📚 SAA 題庫：{len(QUESTIONS['SAA'])} 題",
                                "size": "sm",
                                "color": "#6B6B6B"
                            },
                            {
                                "type": "text",
                                "text": f"📚 SAP 題庫：{len(QUESTIONS['SAP'])} 題",
                                "size": "sm",
                                "color": "#6B6B6B"
                            }
                        ]
                    }
                ]
            },
            "footer": {
                "type": "box",
                "layout": "vertical",
                "spacing": "sm",
                "contents": [
                    {
                        "type": "button",
                        "style": "primary",
                        "color": COLORS['saa'],
                        "action": {
                            "type": "postback",
                            "label": "📘 SAA 助理架構師",
                            "data": "menu=lang&type=SAA"
                        }
                    },
                    {
                        "type": "button",
                        "style": "primary",
                        "color": COLORS['sap'],
                        "action": {
                            "type": "postback",
                            "label": "📗 SAP 專業架構師",
                            "data": "menu=lang&type=SAP"
                        }
                    }
                ]
            }
        }
    }

def lang_select_flex(exam_type):
    """語言選擇"""
    name = "Solutions Architect Associate" if exam_type == "SAA" else "Solutions Architect Professional"
    count = len(QUESTIONS.get(exam_type, []))
    color = COLORS['saa'] if exam_type == "SAA" else COLORS['sap']
    
    return {
        "type": "flex",
        "altText": "選擇語言",
        "contents": {
            "type": "bubble",
            "body": {
                "type": "box",
                "layout": "vertical",
                "contents": [
                    {
                        "type": "text",
                        "text": name,
                        "weight": "bold",
                        "size": "md",
                        "color": color,
                        "wrap": True
                    },
                    {
                        "type": "text",
                        "text": f"題庫共 {count} 題",
                        "size": "sm",
                        "color": "#888888",
                        "margin": "sm"
                    },
                    {
                        "type": "separator",
                        "margin": "md"
                    },
                    {
                        "type": "text",
                        "text": "請選擇答題語言",
                        "margin": "md",
                        "color": "#5A5A5A"
                    }
                ]
            },
            "footer": {
                "type": "box",
                "layout": "vertical",
                "spacing": "sm",
                "contents": [
                    {
                        "type": "button",
                        "style": "primary",
                        "color": COLORS['primary'],
                        "action": {
                            "type": "postback",
                            "label": "🇹🇼 繁體中文",
                            "data": f"action=start&lang=cn&type={exam_type}"
                        }
                    },
                    {
                        "type": "button",
                        "style": "primary",
                        "color": COLORS['accent'],
                        "action": {
                            "type": "postback",
                            "label": "🇺🇸 English",
                            "data": f"action=start&lang=en&type={exam_type}"
                        }
                    }
                ]
            }
        }
    }

def question_flex(exam_type, idx, lang):
    """題目顯示"""
    q = QUESTIONS[exam_type][idx]
    title = q['q_cn'] if lang == 'cn' else q['q_en']
    opts = q['opts_cn'] if lang == 'cn' else q['opts_en']
    
    option_texts = []
    for char in ['A', 'B', 'C', 'D']:
        if opts.get(char):
            option_texts.append({
                "type": "text",
                "text": f"{char}. {opts[char]}",
                "wrap": True,
                "size": "sm",
                "color": "#5A5A5A",
                "margin": "md"
            })
    
    type_color = COLORS['saa'] if exam_type == "SAA" else COLORS['sap']
    
    return {
        "type": "flex",
        "altText": "題目",
        "contents": {
            "type": "bubble",
            "header": {
                "type": "box",
                "layout": "vertical",
                "contents": [{
                    "type": "text",
                    "text": f"{'🇹🇼' if lang == 'cn' else '🇺🇸'} {exam_type} 練習",
                    "size": "xs",
                    "color": "#AAAAAA"
                }]
            },
            "body": {
                "type": "box",
                "layout": "vertical",
                "contents": [
                    {
                        "type": "text",
                        "text": title,
                        "wrap": True,
                        "weight": "bold",
                        "size": "md",
                        "color": type_color
                    },
                    {
                        "type": "separator",
                        "margin": "lg"
                    },
                    {
                        "type": "box",
                        "layout": "vertical",
                        "margin": "lg",
                        "contents": option_texts
                    }
                ]
            },
            "footer": {
                "type": "box",
                "layout": "vertical",
                "contents": [{
                    "type": "box",
                    "layout": "horizontal",
                    "spacing": "sm",
                    "contents": [
                        {
                            "type": "button",
                            "style": "primary",
                            "color": COLORS['secondary'],
                            "action": {
                                "type": "postback",
                                "label": char,
                                "data": f"action=ans&user={char}&correct={q['ans']}&idx={idx}&lang={lang}&type={exam_type}"
                            }
                        } for char in ['A', 'B', 'C', 'D']
                    ]
                }]
            }
        }
    }

@app.route("/")
def home():
    return f"✅ AWS Bot 運行中<br>SAA: {len(QUESTIONS['SAA'])} 題<br>SAP: {len(QUESTIONS['SAP'])} 題"

@app.route("/callback", methods=['POST'])
def callback():
    events = request.json.get('events', [])
    
    for event in events:
        reply_token = event.get('replyToken')
        
        if event['type'] == 'follow' or (event['type'] == 'message' and event['message']['type'] == 'text'):
            send_line(reply_token, [welcome_flex()])
        
        elif event['type'] == 'postback':
            data = parse_qs(event['postback']['data'])
            action = data.get('action', [None])[0]
            lang = data.get('lang', ['cn'])[0]
            exam_type = data.get('type', ['SAA'])[0]
            
            if data.get('menu', [None])[0] == 'lang':
                send_line(reply_token, [lang_select_flex(exam_type)])
            
            elif action == 'start':
                idx = get_random_q(exam_type, lang)
                if idx is not None:
                    send_line(reply_token, [question_flex(exam_type, idx, lang)])
                else:
                    send_line(reply_token, [{"type": "text", "text": f"❌ {exam_type} {'中文' if lang == 'cn' else '英文'}題庫目前沒有題目"}])
            
            elif action == 'ans':
                user_ans = data.get('user', [''])[0]
                correct_ans = data.get('correct', [''])[0]
                q_idx = int(data.get('idx', ['0'])[0])
                
                q = QUESTIONS[exam_type][q_idx]
                explain = q['exp_cn'] if lang == 'cn' else q['exp_en']
                
                result = "🎉 答對了！" if user_ans == correct_ans else f"❌ 答錯了"
                
                send_line(reply_token, [
                    {"type": "text", "text": f"{result}\n{'正確答案' if lang == 'cn' else 'Answer'}: {correct_ans}\n\n💡 {explain}"},
                    {
                        "type": "flex",
                        "altText": "下一步",
                        "contents": {
                            "type": "bubble",
                            "size": "micro",
                            "body": {
                                "type": "box",
                                "layout": "vertical",
                                "spacing": "sm",
                                "contents": [
                                    {
                                        "type": "button",
                                        "style": "primary",
                                        "color": COLORS['primary'],
                                        "action": {
                                            "type": "postback",
                                            "label": "➡️ 下一題",
                                            "data": f"action=start&lang={lang}&type={exam_type}"
                                        }
                                    },
                                    {
                                        "type": "button",
                                        "style": "secondary",
                                        "action": {
                                            "type": "postback",
                                            "label": "🏠 主選單",
                                            "data": "menu=main"
                                        }
                                    }
                                ]
                            }
                        }
                    }
                ])
            
            elif data.get('menu', [None])[0] == 'main':
                send_line(reply_token, [welcome_flex()])
    
    return 'OK'

if __name__ == "__main__":
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
